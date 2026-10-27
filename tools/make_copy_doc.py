"""Collect every piece of site copy into one Word document for Samhan to edit.

Reads site.yaml, the Markdown in content/, projects.yaml, gallery.yaml and notes.yaml, plus the
few strings that live in templates, and writes docs/copy-for-samhan.docx (git-ignored). Each
block is preceded by a grey note saying where it appears and any limit. Samhan edits with track
changes; the accepted text is then transferred back into the source files.

Usage (PowerShell):  python tools/make_copy_doc.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from sitegen.render import load_content  # noqa: E402

OUTPUT = ROOT / "docs" / "copy-for-samhan.docx"
GREY = RGBColor(0x7C, 0x85, 0x8F)

# Strings that live in templates rather than content files, with where they appear.
TEMPLATE_COPY = {
    "Buttons and microcopy (templates)": [
        ("Hero buttons", "Learn more about me / Learn more about my PhD / Reach out"),
        ("About page buttons", "Full CV / My PhD"),
        ("Footer", "Samhan Samhan · Last updated <month year, automatic>"),
        ("Notes back link", "← All notes"),
    ],
    "CV page (templates/cv.html)": [
        ("Kicker", "Full profile"),
        ("Title", "Curriculum vitae"),
        ("View switch", "Featured / Full detail"),
        (
            "Count labels",
            "talks, presentations and publications · certifications · " "awards and scholarships",
        ),
        (
            "Section headings",
            "Education · Professional experience · Professional certifications · "
            "Awards and scholarships · Publications · Conference presentations · "
            "Speaking and media · Volunteering · Professional organisations · "
            "Languages",
        ),
        ("Short-course heading", "Exchanges, internships and short courses"),
    ],
    "Notes pages (templates/notes.html, note.html)": [
        ("Kicker", "Samhan Notes"),
        (
            "Index lede",
            "Posts from my old Blogger site, kept in their original Indonesian with the "
            "photos and the reader comments.",
        ),
        (
            "Post page line",
            "Originally published on Samhan Notes, <month year>, in Indonesian. " "<summary>",
        ),
    ],
    "Gallery page (templates/gallery.html)": [
        ("Kicker", "Moments"),
        ("Title", "Gallery"),
    ],
}


def note(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def add_runs(paragraph, text):
    """Plain text with **bold** spans and [label](url) links rendered as 'label (url)'."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for index, part in enumerate(re.split(r"\*\*", text)):
        if part:
            paragraph.add_run(part).bold = bool(index % 2)


def add_markdown(doc, body):
    """A small Markdown-to-Word mapping: headings, bullets, numbered items, paragraphs."""
    paragraph_lines: list[str] = []

    def flush():
        if paragraph_lines:
            add_runs(doc.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    in_html = False
    for line in body.splitlines():
        stripped = line.strip()
        if stripped.startswith("<div"):
            in_html = True
        if in_html:
            in_html = not stripped.startswith("</div>")
            continue
        marker = re.fullmatch(r"<!--\s*(\w+)\s*-->", stripped)
        if marker:
            flush()
            note(doc, f"[Generated block: {marker.group(1)}. Not editable here.]")
        elif stripped.startswith("### "):
            flush()
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            flush()
            doc.add_heading(stripped[3:], level=2)
        elif re.match(r"^[-*] ", stripped):
            flush()
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            flush()
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", stripped))
        elif stripped == "":
            flush()
        else:
            paragraph_lines.append(stripped)
    flush()


def add_front_matter(doc, meta, where):
    for key in ("kicker", "title", "lede", "description"):
        if meta.get(key):
            note(doc, f"{key.capitalize()} ({where})")
            doc.add_paragraph(str(meta[key]))


def main() -> None:
    settings = yaml.safe_load((ROOT / "site.yaml").read_text(encoding="utf-8"))
    doc = Document()
    doc.add_heading("Website copy for review", level=0)
    doc.add_paragraph(
        "Every piece of text on the site, in page order. Edit with track changes; blocks marked "
        "[Samhan: …] are gaps for you to fill. Grey notes say where each block appears and any "
        "limit. Bold inside a block is formatting on the page, not emphasis to keep. Links are "
        "shown as label (url). Text in the CV page itself comes from the Excel workbook and is "
        "not repeated here."
    )

    doc.add_heading("Home", level=1)
    note(doc, "Hero headline (site.yaml, pronoun-free, one line)")
    doc.add_paragraph(settings["headline"])
    note(doc, "Hero subline (site.yaml, one or two lines)")
    doc.add_paragraph(settings["subline"])
    meta, body = load_content(ROOT / "content" / "home.md")
    note(doc, "In short (content/home.md, at most 300 words, no bold; also the LinkedIn About)")
    add_markdown(doc, body)

    for page, path in (("About", "about.md"), ("PhD", "phd/index.md"), ("Contact", "contact.md")):
        doc.add_heading(page, level=1)
        meta, body = load_content(ROOT / "content" / path)
        add_front_matter(doc, meta, f"content/{path}")
        add_markdown(doc, body)
        if page == "About":
            doc.add_heading("Pet-project cards (content/projects.yaml)", level=2)
            for project in yaml.safe_load((ROOT / "content" / "projects.yaml").read_text("utf-8")):
                note(doc, f"Card: {project['title']} · link label: {project.get('link_label', '')}")
                doc.add_paragraph(project["summary"])

    for heading, items in TEMPLATE_COPY.items():
        doc.add_heading(heading, level=1)
        for label, text in items:
            note(doc, label)
            doc.add_paragraph(text)

    doc.add_heading("Notes summaries (content/notes/notes.yaml)", level=1)
    for entry in yaml.safe_load((ROOT / "content" / "notes" / "notes.yaml").read_text("utf-8")):
        note(doc, f"{entry['title']} ({entry['date']}) · picture: {entry.get('image', '')}")
        doc.add_paragraph(entry["summary"])

    doc.add_heading("Gallery captions (content/gallery.yaml)", level=1)
    for item in yaml.safe_load((ROOT / "content" / "gallery.yaml").read_text("utf-8")):
        note(doc, f"{item['file']} · {item.get('year', '')}")
        doc.add_paragraph(item["caption"])

    doc.add_heading("Appendix: linked organisations (site.yaml org_links)", level=1)
    doc.add_paragraph(
        "These names become links wherever they appear. Confirm each address; the NSW "
        "Decarbonisation Innovation Hub and Digital GreenTalents ones are unverified."
    )
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Name", "URL"
    for name, url in settings["org_links"].items():
        row = table.add_row().cells
        row[0].text, row[1].text = name, url

    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print("wrote", OUTPUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
